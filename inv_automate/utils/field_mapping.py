"""Parse client field-mapping Excel (.xlsx) into text for prompt generation."""

from __future__ import annotations

import io
import re
from collections import defaultdict
from typing import Any

# Exact headers expected in the UI / docs
EXPECTED_HEADERS = (
    "Table Name",
    "Column Name",
    "Field Location in Actual Invoice",
    "Remarks",
)

_HEADER_ALIASES = {
    "table name": "table",
    "table": "table",
    "column name": "column",
    "column": "column",
    "field name": "column",
    "field location in actual invoice": "location",
    "field location": "location",
    "location": "location",
    "remarks": "remarks",
    "remark": "remarks",
    "examples": "remarks",
    "example": "remarks",
}


def _norm_header(cell: Any) -> str:
    text = str(cell or "").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_field_mapping_xlsx(file_bytes: bytes) -> tuple[str, int, list[dict[str, str]]]:
    """
    Parse a single-sheet field mapping workbook.

    Returns:
        (formatted_text, row_count, rows)
        rows: [{"table", "column", "location", "remarks"}, ...]

    Raises:
        ValueError with a user-facing message on bad input.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            "openpyxl is not installed. Run: pip install openpyxl"
        ) from exc

    if not file_bytes:
        raise ValueError("Empty field mapping file.")

    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Could not read Excel file: {exc}") from exc

    if len(wb.sheetnames) != 1:
        raise ValueError(
            f"Field mapping workbook must contain exactly one sheet "
            f"(found {len(wb.sheetnames)}). Keep only one sheet and retry."
        )

    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Field mapping sheet is empty.")

    header_idx = None
    col_map: dict[str, int] = {}
    for i, row in enumerate(rows[:25]):
        cells = [_norm_header(c) for c in row]
        found: dict[str, int] = {}
        for j, cell in enumerate(cells):
            key = _HEADER_ALIASES.get(cell)
            if key and key not in found:
                found[key] = j
        if "column" in found and "location" in found:
            # remarks optional in header match but preferred
            header_idx = i
            col_map = found
            break

    if header_idx is None or "column" not in col_map:
        raise ValueError(
            "Could not find header row. First row (or near top) must include: "
            + ", ".join(EXPECTED_HEADERS)
        )

    if "remarks" not in col_map:
        # Allow missing Remarks column but warn via empty remarks
        pass

    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    flat_rows: list[dict[str, str]] = []
    kept = 0
    for row in rows[header_idx + 1 :]:
        if not row:
            continue
        table = _cell_str(row[col_map["table"]]) if "table" in col_map and col_map["table"] < len(row) else ""
        column = _cell_str(row[col_map["column"]]) if col_map["column"] < len(row) else ""
        location = (
            _cell_str(row[col_map["location"]])
            if "location" in col_map and col_map["location"] < len(row)
            else ""
        )
        remarks = (
            _cell_str(row[col_map["remarks"]])
            if "remarks" in col_map and col_map["remarks"] < len(row)
            else ""
        )

        if not column:
            continue
        if not location and not remarks:
            continue

        table_key = table or "(unspecified table)"
        item = {
            "table": table_key,
            "column": column,
            "location": location,
            "remarks": remarks,
        }
        grouped[table_key].append(item)
        flat_rows.append(item)
        kept += 1

    if kept == 0:
        raise ValueError(
            "No usable mapping rows found. Each kept row needs a Column Name "
            "and at least a Location or Remarks value."
        )

    lines = [
        "FIELD MAPPING (locations and remarks by table).",
        "JSON Schema defines the output fields; use this sheet only as location/remarks hints.",
        "If a schema field is missing here, or Remarks is empty, do not invent details.",
        "Treat 'Default: …' in Location as an extraction default rule.",
        "",
    ]
    for table_name in grouped:
        lines.append(f"[{table_name}]")
        for item in grouped[table_name]:
            lines.append(f"- {item['column']}")
            if item["location"]:
                lines.append(f"  location: {item['location']}")
            if item["remarks"]:
                lines.append(f"  remarks: {item['remarks']}")
            else:
                lines.append("  remarks: (none)")
        lines.append("")

    return "\n".join(lines).strip(), kept, flat_rows
